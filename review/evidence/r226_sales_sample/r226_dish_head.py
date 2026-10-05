# -*- coding: utf-8 -*-
import openpyxl, os
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
f = [x for x in os.listdir(D) if '菜品销售统计' in x][0]
wb = openpyxl.load_workbook(os.path.join(D, f), data_only=True)
ws = wb.worksheets[0]
print('sheet =', ws.title, ws.dimensions)
for r in range(1, 34):
    row = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        row.append('' if v is None else str(v))
    print('R%-4d| %s' % (r, ' | '.join(row)))
print('...')
for r in range(ws.max_row - 6, ws.max_row + 1):
    row = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        row.append('' if v is None else str(v))
    print('R%-4d| %s' % (r, ' | '.join(row)))
