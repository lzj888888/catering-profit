# -*- coding: utf-8 -*-
import openpyxl, os
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
def wbof(kw):
    f = [x for x in os.listdir(D) if kw in x and x.endswith('.xlsx')][0]
    return openpyxl.load_workbook(os.path.join(D, f), data_only=True)

print('=' * 84)
print('【A】饿了么门店日表 —— 平台等式逐行校验（7 天）')
ws = wbof('门店下载').worksheets[0]
hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}
rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)] for r in range(2, ws.max_row + 1)]
def g(row, name): 
    v = row[ix[name] - 1] if name in ix else None
    return 0 if v is None else v
ok = True
for row in rows:
    lhs = g(row, '收入')
    rev = g(row, '营业额')
    exp = g(row, '支出')
    a = round(rev - exp, 2)
    b = round(g(row,'商品销售额')+g(row,'打包费')+g(row,'商家应收配送费')+g(row,'索赔单')+g(row,'其他营业额'), 2)
    c = round(sum(g(row,k) for k in ['平台技术服务费','活动补贴','代金券补贴','配送费补贴','智能满减补贴','智能满减服务费','履约技术服务费','退单费用','其他支出']), 2)
    flag = (abs(lhs-a) < 0.02) and (abs(rev-b) < 0.02) and (abs(exp-c) < 0.02)
    ok &= flag
    if not flag: print('   ✗', row[0], lhs, a, rev, b, exp, c)
print('   收入=营业额-支出 → 7/7 成立 =', ok)
tot = {k: round(sum(g(r,k) for r in rows), 2) for k in ['有效订单','无效订单','收入','营业额','商品销售额','打包费','支出','平台技术服务费','活动补贴','代金券补贴','配送费补贴','智能满减补贴','智能满减服务费','履约技术服务费','顾客实付总额','活动总补贴','饿了么补贴','代理商补贴','商家活动成本（含满减活动）']}
print('   7 天合计：', {k: v for k, v in tot.items() if k in ['有效订单','收入','营业额','商品销售额','打包费','支出','顾客实付总额']})
print('   费用明细：', {k: v for k, v in tot.items() if k in ['平台技术服务费','活动补贴','代金券补贴','配送费补贴','智能满减补贴','智能满减服务费','履约技术服务费']})
print('   补贴：', {k: v for k, v in tot.items() if k in ['活动总补贴','饿了么补贴','代理商补贴','商家活动成本（含满减活动）']})
print('   到手率 = 收入/营业额 = %.2f%%' % (tot['收入']/tot['营业额']*100))
print('   平台抽成(技术+履约)/营业额 = %.2f%%' % ((tot['平台技术服务费']+tot['履约技术服务费'])/tot['营业额']*100))

print()
print('=' * 84)
print('【B】POS 菜品销量表 —— 合计与结构')
ws2 = wbof('菜品销售统计').worksheets[0]
dish = {}
for r in range(5, ws2.max_row):
    n = ws2.cell(row=r, column=1).value
    if n is None: continue
    dish[n] = [ws2.cell(row=r, column=c).value for c in range(2, 11)]
totrow = [ws2.cell(row=ws2.max_row, column=c).value for c in range(2, 11)]
print('   行数(不含合计) =', len(dish), ' 合计行 =', totrow)
s_qty = sum(v[0] for v in dish.values()); s_amt = sum(v[2] for v in dish.values())
s_inc = sum(v[4] for v in dish.values()); s_dis = sum(v[6] for v in dish.values())
print('   逐行求和: 数量=%.1f 销售额=%.1f 收入=%.1f 优惠=%.1f' % (s_qty, s_amt, s_inc, s_dis))
print('   合计行对账: 数量 %s  销售额 %s  收入 %s  优惠 %s' % (totrow[0], totrow[2], totrow[4], totrow[6]))
print('   收入 = 销售额 - 优惠 → %.2f vs %.2f' % (s_amt - s_dis, s_inc))
top = sorted(dish.items(), key=lambda kv: -(kv[1][2] or 0))[:15]
print('   TOP15（按销售额）:')
for n, v in top:
    print('     %-26s 数量=%-8s 销售额=%-9s 占比=%.2f%%' % (n, v[0], v[2], (v[3] or 0)*100))

print()
print('=' * 84)
print('【C】POS 套餐明细表 —— 合计与套餐清单')
ws3 = wbof('套餐销售统计').worksheets[0]
combos = {}
for r in range(4, ws3.max_row):
    cn = ws3.cell(row=r, column=2).value
    if cn is None: continue
    v = [ws3.cell(row=r, column=c).value for c in (6, 7, 8, 9, 10)]
    com = combos.setdefault(cn, {'cnt': 0, 'qty': 0.0, 'amt': 0.0, 'dis': 0.0, 'inc': 0.0, 'items': 0, 'big': set()})
    com['cnt'] += 1; com['qty'] += v[1] or 0; com['amt'] += v[2] or 0
    com['dis'] += v[3] or 0; com['inc'] += v[4] or 0; com['big'].add(v[0])
print('   套餐数 =', len(combos), ' 明细行 =', sum(c['cnt'] for c in combos.values()))
tq = sum(c['qty'] for c in combos.values()); ta = sum(c['amt'] for c in combos.values())
ti = sum(c['inc'] for c in combos.values()); td = sum(c['dis'] for c in combos.values())
print('   合计: 数量=%.1f 销售额=%.1f 收入=%.1f 优惠=%.1f' % (tq, ta, ti, td))
for cn, c in sorted(combos.items(), key=lambda kv: -kv[1]['amt']):
    print('     %-56s 行=%-3d 数量=%-6.1f 销售额=%-8.1f 收入=%-8.1f 大类=%s' % (cn[:54], c['cnt'], c['qty'], c['amt'], c['inc'], '/'.join(sorted(c['big']))))

print()
print('=' * 84)
print('【D】🔴 重复计算判据：套餐明细里的单品，是否已出现在菜品表里')
comp = set()
for r in range(4, ws3.max_row):
    n = ws3.cell(row=r, column=3).value
    if n: comp.add(n)
hit = sorted(n for n in comp if n in dish)
print('   套餐组分单品名 =', len(comp), ' 其中在菜品表命中 =', len(hit))
print('   命中的 =', hit)
allhit = sorted(n for n in comp if n in dish)
miss = sorted(n for n in comp if n not in dish)
print('   未命中的 =', miss)
