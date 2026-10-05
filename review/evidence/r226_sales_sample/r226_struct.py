# -*- coding: utf-8 -*-
import openpyxl, os
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
def tof(v):
    if v is None or v == '': return 0.0
    try: return float(v)
    except Exception: return 0.0
wb = openpyxl.load_workbook(os.path.join(D, [x for x in os.listdir(D) if '菜品销售统计' in x][0]), data_only=True)
ws = wb.worksheets[0]
items = []
for r in range(5, ws.max_row):
    n = ws.cell(row=r, column=1).value
    if n is None: continue
    items.append((n, tof(ws.cell(row=r, column=2).value), tof(ws.cell(row=r, column=4).value), tof(ws.cell(row=r, column=6).value), tof(ws.cell(row=r, column=8).value)))
TOT_A = sum(x[2] for x in items); TOT_I = sum(x[3] for x in items)
print('堂食 菜品行 = %d  销售额 = %.1f  收入 = %.1f' % (len(items), TOT_A, TOT_I))

CIG = ['香烟', '龙凤呈祥', '云烟', '中华', '玉溪', '芙蓉王', '利群']
DRK = ['东鹏', '农夫山泉', '百事', '可乐', '雪碧', '康师傅', '统一', '阿萨姆', '乐堡', '国宾', '啤酒', '冰红茶', '茉莉', '矿泉水', '苏打水', '脉动', '红牛', '娃哈哈', '尖叫', '元气森林']
PKG = ['打包盒', '餐盒']
def bucket(n):
    if any(k in n for k in CIG): return 'C 烟草'
    if any(k in n for k in PKG): return 'D 包装耗材'
    if any(k in n for k in DRK): return 'B 酒水饮料'
    return 'A 厨房出品'
agg = {}
for n, q, a, i, d in items:
    b = bucket(n)
    x = agg.setdefault(b, [0, 0.0, 0.0])
    x[0] += 1; x[1] += a; x[2] += i
print()
print('%-12s %5s %12s %8s %12s %8s' % ('桶', 'SKU', '销售额', '占比', '收入', '占比'))
for b in sorted(agg):
    x = agg[b]
    print('%-12s %5d %12.1f %7.2f%% %12.1f %7.2f%%' % (b, x[0], x[1], x[1]/TOT_A*100, x[2], x[2]/TOT_I*100))
nk = agg['A 厨房出品']; print()
print('⇒ 厨房出品（真餐饮）销售额 %.1f = %.2f%%；非厨房（烟+酒饮+包装）%.1f = %.2f%%'
      % (nk[1], nk[1]/TOT_A*100, TOT_A-nk[1], (TOT_A-nk[1])/TOT_A*100))

print()
print('--- 快餐/面食线（15元快餐+5元+小面+韭叶+米线+抄手）---')
LINE = ['快餐15元', '快餐5元', '小面', '韭叶', '米线', '抄手', '刀削面', '酸辣粉']
s = 0.0; q = 0.0
for n, qq, a, i, d in items:
    if any(n == k or n.startswith(k) for k in LINE):
        s += a; q += qq; print('   %-16s 数量=%-8.1f 销售额=%-9.1f 收入=%.1f' % (n, qq, a, i))
print('   小计 数量=%.1f 销售额=%.1f = %.2f%% of 堂食' % (q, s, s/TOT_A*100))

print()
print('--- 火锅线（耙牛肉/耙牛筋/锅底/小料等）---')
HOT = ['耙牛肉', '耙牛筋', '耙肥肠', '红汤', '鸳鸯', '小料', '自助小料', '毛肚', '鸭肠', '黄喉', '牛肉滑', '酥肉']
s2 = 0.0
for n, qq, a, i, d in items:
    if any(k in n for k in HOT):
        s2 += a
        if a > 60: print('   %-20s 数量=%-8.1f 销售额=%-9.1f' % (n, qq, a))
print('   火锅相关合计（堂食，含套餐拆解）销售额 = %.1f = %.2f%% of 堂食' % (s2, s2/TOT_A*100))

print()
print('--- 外卖（饿了么 9/7–9/13）---')
ws3 = openpyxl.load_workbook(os.path.join(D, [x for x in os.listdir(D) if '门店下载' in x and '(1)' not in x][0]), data_only=True).worksheets[0]
hdr = [ws3.cell(row=1, column=c).value for c in range(1, ws3.max_column + 1)]
ix = {h: i for i, h in enumerate(hdr) if h}
S = lambda r, k: tof(ws3.cell(row=r, column=ix[k] + 1).value)
rows = range(2, ws3.max_row + 1)
wr = sum(S(r, '营业额') for r in rows); wrc = sum(S(r, '收入') for r in rows)
print('   营业额 %.2f  到手 %.2f  有效订单 %d' % (wr, wrc, sum(S(r, '有效订单') for r in rows)))
print('   堂食(销售额) %.1f  vs  外卖(营业额) %.1f  ⇒ 外卖占比 %.2f%%' % (TOT_A, wr, wr/(TOT_A+wr)*100))
print('   堂食(收入)   %.1f  vs  外卖(到手)   %.1f  ⇒ 外卖占比 %.2f%%' % (TOT_I, wrc, wrc/(TOT_I+wrc)*100))
print()
print('--- 活动总补贴 口径核（逐行）---')
bad = 0
for r in rows:
    a = S(r, '活动总补贴'); b = S(r, '饿了么补贴') + S(r, '代理商补贴') + S(r, '商家活动成本（含满减活动）')
    if abs(a - b) > 0.02:
        bad += 1
        print('   ✗ %s  活动总补贴=%.2f  分项和=%.2f  差=%.2f' % (ws3.cell(row=r, column=1).value, a, b, a - b))
print('   7 天中不符行数 =', bad, '/ 7')
print()
print('--- 套餐重复计算幅度 ---')
ws2 = openpyxl.load_workbook(os.path.join(D, [x for x in os.listdir(D) if '套餐销售统计' in x][0]), data_only=True).worksheets[0]
ca = sum(tof(ws2.cell(row=r, column=8).value) for r in range(4, ws2.max_row))
print('   套餐明细销售额 = %.1f ；堂食总销售额 = %.1f ⇒ 套餐占堂食 %.2f%%' % (ca, TOT_A, ca/TOT_A*100))
print('   若"菜品表 + 套餐表"直接相加 ⇒ 虚增 %.1f 元（%.2f%%）' % (ca, ca/(TOT_A+ca)*100))
