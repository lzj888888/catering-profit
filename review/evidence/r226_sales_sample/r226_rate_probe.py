# -*- coding: utf-8 -*-
"""R226：用 7 天真实数据判别「平台技术服务费」的计提基数（含/不含打包费）"""
import openpyxl, os
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
ws = openpyxl.load_workbook(os.path.join(D, [x for x in os.listdir(D) if '门店下载' in x and '(1)' not in x][0]), data_only=True).worksheets[0]
hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}
def S(r, k): 
    v = ws.cell(row=r, column=ix[k]).value
    try: return float(v)
    except Exception: return 0.0
def stat(xs):
    return min(xs), max(xs), sum(xs) / len(xs), max(xs) - min(xs)
print('%-12s %10s %10s %10s %10s' % ('日期', '技术费', '商品销售额', '营业额', '费率(商品)'))
r_goods, r_rev = [], []
for r in range(2, ws.max_row + 1):
    d = ws.cell(row=r, column=1).value
    fee = S(r, '平台技术服务费'); g = S(r, '商品销售额'); rev = S(r, '营业额')
    a, b = fee / g * 100, fee / rev * 100
    r_goods.append(a); r_rev.append(b)
    print('%-12s %10.2f %10.2f %10.2f %9.4f%%  (÷营业额 %7.4f%%)' % (d, fee, g, rev, a, b))
for nm, xs in [('÷ 商品销售额（不含打包费）', r_goods), ('÷ 营业额（含打包费）', r_rev)]:
    mn, mx, avg, sp = stat(xs)
    print('%-26s min=%.4f%% max=%.4f%% avg=%.4f%% 极差=%.4fpp' % (nm, mn, mx, avg, sp))
print()
print('⇒ 判定：极差更小的一侧更可能是计提基数（打包费仅占营业额 2.52%，故信号弱，仅供参考）')
print()
# 履约技术服务费 同测
print('--- 履约技术服务费 基数 ---')
g1, g2 = [], []
for r in range(2, ws.max_row + 1):
    f = S(r, '履约技术服务费'); g = S(r, '商品销售额'); rev = S(r, '营业额')
    if rev: g1.append(f/g*100); g2.append(f/rev*100)
print('÷ 商品销售额 avg=%.4f%% 极差=%.4fpp ; ÷ 营业额 avg=%.4f%% 极差=%.4fpp'
      % (sum(g1)/len(g1), max(g1)-min(g1), sum(g2)/len(g2), max(g2)-min(g2)))
print()
print('--- 合计口径 ---')
TF = sum(S(r, '平台技术服务费') for r in range(2, ws.max_row+1))
PF = sum(S(r, '履约技术服务费') for r in range(2, ws.max_row+1))
G = sum(S(r, '商品销售额') for r in range(2, ws.max_row+1))
R = sum(S(r, '营业额') for r in range(2, ws.max_row+1))
print('技术费 %.2f / 商品 %.2f = %.4f%% ; / 营业额 %.2f = %.4f%%' % (TF, G, TF/G*100, R, TF/R*100))
print('履约费 %.2f / 商品 %.2f = %.4f%% ; / 营业额 %.2f = %.4f%%' % (PF, G, PF/G*100, R, PF/R*100))
print('两项合计 %.2f = 营业额的 %.4f%%' % (TF+PF, (TF+PF)/R*100))
