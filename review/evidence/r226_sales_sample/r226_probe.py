# -*- coding: utf-8 -*-
import openpyxl, os, hashlib, sys
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
files = sorted(f for f in os.listdir(D) if f.lower().endswith('.xlsx'))
print("目录内 xlsx 数 =", len(files))
for f in files:
    p = os.path.join(D, f)
    b = open(p, 'rb').read()
    print("=" * 90)
    print("FILE:", f)
    print("  size = %d B   md5 = %s" % (len(b), hashlib.md5(b).hexdigest()))
    wb = openpyxl.load_workbook(p, data_only=True)
    for ws in wb.worksheets:
        print("  [sheet] %-24s dims=%-12s max_row=%d max_col=%d" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
