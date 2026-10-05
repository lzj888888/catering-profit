# -*- coding: utf-8 -*-
import openpyxl, os
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
f = [x for x in os.listdir(D) if '门店下载' in x and '(1)' not in x][0]
ws = openpyxl.load_workbook(os.path.join(D, f), data_only=True).worksheets[0]
hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print('表头重复项:')
from collections import Counter
c = Counter([h for h in hdr if h])
print('  ', [k for k, v in c.items() if v > 1])
print('非数值单元格（前 20）:')
n = 0
for r in range(2, ws.max_row + 1):
    for ci in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=ci).value
        if isinstance(v, str):
            n += 1
            if n <= 20:
                print('   R%d C%d (%s) = %r' % (r, ci, hdr[ci-1], v))
print('非数值总数 =', n)
