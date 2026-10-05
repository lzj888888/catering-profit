# -*- coding: utf-8 -*-
import openpyxl, os
D = r"C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/temp/RWTemp/2026-10/9c5c48749ff36941714d28465cf0bce5"
OUT = r"C:/Users/lzj/WorkBuddy/2026-09-08-22-08-11/_r226_data"
tagmap = {
 '套餐销售统计': 'combo',
 '菜品销售统计': 'dish',
 '门店下载_20260907至20260913__5366951484_20261005190640294.xlsx': 'store_a',
}
for f in sorted(os.listdir(D)):
    if not f.lower().endswith('.xlsx'): continue
    tag = None
    for k, v in tagmap.items():
        if k in f and not (k == '门店下载' and '(1)' in f):
            tag = v
    if tag is None: continue
    p = os.path.join(D, f)
    wb = openpyxl.load_workbook(p, data_only=True)
    lines = ['FILE: ' + f, '']
    for ws in wb.worksheets:
        lines.append('#### sheet: %s  (%s)  max_row=%d max_col=%d' % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
        for r in range(1, ws.max_row + 1):
            row = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                row.append('' if v is None else str(v))
            lines.append('R%-4d| ' % r + ' | '.join(row))
        lines.append('')
    op = os.path.join(OUT, tag + '.txt')
    open(op, 'w', encoding='utf-8', newline='\n').write('\n'.join(lines))
    print('written', tag, os.path.getsize(op), 'chars', len('\n'.join(lines)))
