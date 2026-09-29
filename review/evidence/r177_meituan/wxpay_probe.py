# -*- coding: utf-8 -*-
"""R178b：确认 ① 行数 97/98 差异 ② 文件名写「微信支付」但内容属哪个平台"""
import sys, io
from collections import defaultdict, Counter
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import openpyxl
P = r'C:\Users\lzj\Desktop\耙三样_耙牛肉2026-08-01_2026-08-31微信支付账单.xlsx'

wb = openpyxl.load_workbook(P, data_only=True)   # 非 read_only，拿真实维度
print('=== 真实维度（非 read_only）===')
for ws in wb.worksheets:
    print(' %-10s max_row=%d max_col=%d' % (ws.title, ws.max_row, ws.max_column))

ws = wb['订单明细']
rows = list(ws.iter_rows(values_only=True))
print()
print('订单明细 iter_rows 行数 =', len(rows))
# 找空行
empty = [i + 1 for i, r in enumerate(rows) if all(v in (None, '') for v in r)]
print('全空行位置:', empty[:20], '共', len(empty))
# 尾部
for i in range(max(0, len(rows) - 4), len(rows)):
    r = rows[i]
    nz = [x for x in r if x not in (None, '')]
    print('  行%3d 非空%d个: %s' % (i + 1, len(nz), str(nz[:6])[:110]))

hdr = [str(h).strip() if h is not None else '' for h in rows[0]]
def ci(name):
    return hdr.index(name) if name in hdr else None

data = rows[1:]
i_type = ci('交易类型'); i_net = ci('商家应收款'); i_pay = ci('用户支付方式')
i_desc = ci('交易描述'); i_oid = ci('订单号'); i_sid = ci('结算id')
print()
print('=== 平台归属判定 ===')
print(' 是否含美团专属列:', [c for c in hdr if '美团' in c or '配送服务费' in c or '佣金' in c][:8])
print(' 用户支付方式分布:')
for k, v in Counter(str(r[i_pay]).strip() for r in data if i_pay is not None).most_common(10):
    print('   %-20s %d' % (k, v))

def num(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(',', '').replace('元', '').strip()
    if s in ('', '-', '--'): return 0.0
    try: return float(s)
    except Exception: return 0.0

print()
print('=== 交易类型 × 金额（与 R177 锚点对表）===')
d = defaultdict(lambda: [0, 0.0])
for r in data:
    t = str(r[i_type]).strip() if i_type is not None else ''
    d[t][0] += 1
    d[t][1] += num(r[i_net])
tot_r = 0; tot_a = 0.0
for t in sorted(d, key=lambda x: -abs(d[x][1])):
    print(' %-30s %4d 行  %10.2f' % (t, d[t][0], d[t][1]))
    tot_r += d[t][0]; tot_a += d[t][1]
print(' %-30s %4d 行  %10.2f' % ('合计', tot_r, round(tot_a, 2)))

print()
print('=== 关键锚点复算 ===')
for name in ['商品总价', '打包费', '商家对顾客的活动补贴', '平台服务费(含佣金和配送服务费)',
             '佣金(实付)', '配送服务费(含补贴)', '商家应收款']:
    i = ci(name)
    if i is None:
        print(' %-34s (列缺失)' % name); continue
    print(' %-34s %10.2f' % (name, round(sum(num(r[i]) for r in data), 2)))

print()
print('=== 汇总信息（原始）===')
ws3 = wb['汇总信息']
for r in ws3.iter_rows(values_only=True):
    nz = [x for x in r if x not in (None, '')]
    if nz: print('  ', nz)

print()
print('=== 账单明细首行/末行 ===')
ws1 = wb['账单明细']
r1 = list(ws1.iter_rows(values_only=True))
for r in r1[:3]:
    nz = [x for x in r if x not in (None, '')]
    if nz: print('  ', str(nz)[:150])
print('   ...')
for r in r1[-2:]:
    nz = [x for x in r if x not in (None, '')]
    if nz: print('  ', str(nz)[:150])
