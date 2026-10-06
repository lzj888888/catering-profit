#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
R232 · 形态 C（外卖/商品级销量）样例机核 dump —— 可复现脚本
=========================================================================
输入（同目录）：
  商品下载_20260907至20260913__5366951484_20261006213915236.xlsx
    md5 = ba724ef6d03dafd478bced4948c1cbe1
    size = 36740 bytes
来源：2026-10-06 21:39 李老师微信发来（发件人「·晨曦」），
      文件名时间戳 20261006213915236。

运行：
  C:/Users/lzj/.workbuddy/binaries/python/envs/default/Scripts/python.exe dump_formc.py

输出：dump_formc.out.txt（stdout 重定向即可）
=========================================================================
"""
import os
import sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit('需要 openpyxl：pip install openpyxl')

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, '商品下载_20260907至20260913__5366951484_20261006213915236.xlsx')


def num(v):
    if v is None or v == '':
        return None
    try:
        return float(v)
    except Exception:
        return None


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    print('== 文件 ================================================')
    print('path =', os.path.basename(XLSX))
    print('size =', os.path.getsize(XLSX), 'bytes')
    print('sheets =', wb.sheetnames)

    print()
    print('== sheet 结构 ==========================================')
    for ws in wb.worksheets:
        print('  %-8s dims=%-12s rows=%-5d cols=%d' % (
            ws.title, ws.dimensions, ws.max_row, ws.max_column))

    ws = wb['data']
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    data = rows[1:]

    print()
    print('== 表头（单行，R1）====================================')
    for i, h in enumerate(hdr, 1):
        print('  %2d  %s' % (i, h))

    print()
    print('== meta sheet =========================================')
    for r in wb['meta'].iter_rows(values_only=True):
        print('  ', r)

    print()
    print('== 规模 ===============================================')
    dates = sorted({r[0] for r in data})
    dishes = sorted({r[3] for r in data})
    shops = sorted({(r[2], r[1]) for r in data})
    print('  data 行数（不含表头） =', len(data))
    print('  日期数 =', len(dates), dates)
    print('  门店数 =', len(shops), shops)
    print('  商品数（去重） =', len(dishes))
    print('  网格完整性: %d 商品 x %d 日 = %d  实得 %d  缺 %d' % (
        len(dishes), len(dates), len(dishes) * len(dates), len(data),
        len(dishes) * len(dates) - len(data)))

    per = defaultdict(set)
    for r in data:
        per[r[3]].add(r[0])
    miss = {d: sorted(set(dates) - v) for d, v in per.items() if set(dates) - v}
    print('  缺日商品数 =', len(miss))
    for d, m in miss.items():
        print('     %-20s 缺 %s' % (d, m))

    print()
    print('== 列类型直方图 =======================================')
    for j, h in enumerate(hdr):
        c = Counter(type(r[j]).__name__ for r in data)
        print('  %2d %-14s %s' % (j + 1, h, dict(c)))

    print()
    print('== 标记列取值 =========================================')
    print('  是否配料:', dict(Counter(r[4] for r in data)))
    print('  是否售罄:', dict(Counter(r[5] for r in data)))

    print()
    print('== 空值统计 ===========================================')
    anynull = False
    for j, h in enumerate(hdr):
        n = sum(1 for r in data if r[j] is None or r[j] == '')
        if n:
            anynull = True
            print('  %-14s 空值 %d' % (h, n))
    if not anynull:
        print('  无空值')

    print()
    print('== 聚合（口径候选）=====================================')
    t_s = sum(num(r[6]) or 0 for r in data)
    t_q = sum(num(r[7]) or 0 for r in data)
    print('  SUM 销售额 = %.4f' % t_s)
    print('  SUM 销量   = %.4f' % t_q)
    nz = [r for r in data if (num(r[7]) or 0) > 0]
    print('  有销量行数 = %d' % len(nz))
    print('  SUM 销量(有销量行) = %.4f' % (sum(num(r[7]) or 0 for r in nz)))
    print('  SUM 销售额(有销量行) = %.4f' % (sum(num(r[6]) or 0 for r in nz)))

    print()
    print('== 逐菜品周汇总（按 qty 降序）==========================')
    agg = defaultdict(lambda: [0.0, 0.0])
    for r in data:
        agg[r[3]][0] += num(r[7]) or 0
        agg[r[3]][1] += num(r[6]) or 0
    lst = sorted(agg.items(), key=lambda kv: -kv[1][0])
    print('  %-52s %8s %12s %10s' % ('商品', 'qty', '销售额', '均价'))
    for d, (q, s) in lst:
        u = ('%.4f' % (s / q)) if q else '-'
        print('  %-52s %8.1f %12.2f %10s' % (d[:52], q, s, u))
    print('  有销量商品数 = %d' % sum(1 for _, (q, _s) in lst if q > 0))
    print('  TOTAL qty = %.1f   TOTAL 销售额 = %.2f' % (
        sum(v[0] for v in agg.values()), sum(v[1] for v in agg.values())))

    print()
    print('== 异常探测 ===========================================')
    zp = defaultdict(lambda: [0.0, 0])
    for r in data:
        q = num(r[7]) or 0
        s = num(r[6]) or 0
        if q > 0 and s == 0:
            zp[r[3]][0] += q
            zp[r[3]][1] += 1
    print('  「有销量但销售额=0」的 SKU:')
    if zp:
        for k, (q, n) in zp.items():
            print('     %-20s qty=%-7.1f 行数=%d' % (k, q, n))
    else:
        print('     无')

    nonint = [(r[3], r[0], num(r[7])) for r in data if (num(r[7]) or 0) != int(num(r[7]) or 0)]
    print('  非整数 qty 行数 =', len(nonint))
    for d, dt, q in nonint[:10]:
        print('     %-20s %s qty=%s' % (d, dt, q))

    print()
    print('== 销售额 vs 订单交易额 ================================')
    mismatch = 0
    for r in data:
        q = num(r[7]) or 0
        s = num(r[6]) or 0
        t = num(r[10]) or 0
        if q > 0 and abs(s - t) > 1e-6:
            mismatch += 1
    print('  同一行 销售额 != 订单交易额 的行数 =', mismatch, '/', len(nz))
    print('  ⇒ 两者语义不同（销售额=菜品口径；订单交易额=整单口径，含其他菜品）')


if __name__ == '__main__':
    main()
