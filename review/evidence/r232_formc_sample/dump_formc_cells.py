#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""R232 · 形态 C 样例 —— **逐格** dump（与 r226_sales_sample/dish.txt 同格式，供 Node 锚点直接 parseDump）

为什么另建一份：既有 dump_formc.out.txt 是**汇总格式**（表头清单 + 聚合值），
喂不进 detectDishShape(sheetRows)——它要的是 0-based 二维 cells。
锚点纪律（spec-increment-authoring）：断言不许抄 NOTE 的二手结论，要从**原始素材**自己解析。

运行：
  C:/Users/lzj/.workbuddy/binaries/python/envs/default/Scripts/python.exe dump_formc_cells.py
输出：formc_cells.txt
"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "商品下载_20260907至20260913__5366951484_20261006213915236.xlsx")

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["data"]

lines = []
lines.append("FILE: " + os.path.basename(SRC))
lines.append("")
lines.append("#### sheet: data  (A1:R%d)  max_row=%d max_col=%d" % (ws.max_row, ws.max_row, ws.max_column))
for r in range(1, ws.max_row + 1):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        cells.append("" if v is None else str(v))
    lines.append("R%d   | %s | " % (r, " | ".join(cells)))
# 🔴 只 dump data sheet：锚点的 parseDump 不区分 sheet，若同时输出 meta（A1:B3），
#    meta 的 R1~R3 会**覆盖** data 的表头行 ⇒ 首跑算出「末日期 = '日期'」、行数多 2（我方构造错）。
#    meta 的结构信息已在既有 dump_formc.out.txt 汇总件里，逐格件专注 data 即可。

out = os.path.join(HERE, "formc_cells.txt")
with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written:", out, "rows:", ws.max_row, "cols:", ws.max_column)
