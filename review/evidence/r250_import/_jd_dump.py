# -*- coding: utf-8 -*-
"""R249-2 诊断：dump 京东两份对账单的真实结构（sheet 名 / 维度 / 前 8 行）"""
import io, sys
from openpyxl import load_workbook

BASE = r'C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/msg/file/2026-10/'
FILES = [
    '208386489_对账单下载_20260930_耙三样耙牛肉_7647722_7647723_0.xlsx',
    '208386489_sku对账单下载_20260930_耙三样耙牛肉_7647724_7647725_0.xlsx',
]
out = io.open(r'C:/Users/lzj/WorkBuddy/Claw/catering-profit/_probe_tmp/_jd_dump.txt', 'w', encoding='utf-8')
def P(s=''):
    print(s); out.write(str(s) + '\n')

for f in FILES:
    p = BASE + f
    P('=' * 100)
    P('文件 = %s' % f)
    try:
        wb = load_workbook(p, data_only=True, read_only=True)
    except Exception as ex:
        P('  打不开：%s' % ex); continue
    P('  sheet 列表 = %r' % (wb.sheetnames,))
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        P('  ---- sheet %r : %d 行 × %d 列' % (sn, len(rows), ws.max_column or 0))
        for i, r in enumerate(rows[:8]):
            cells = ['' if c is None else str(c).strip() for c in r]
            # 截断显示
            P('     R%-2d %s' % (i + 1, cells[:14]))
        if len(rows) > 8:
            P('     ... 其余 %d 行' % (len(rows) - 8))
    wb.close()
out.close()
print('已落盘 _probe_tmp/_jd_dump.txt')
