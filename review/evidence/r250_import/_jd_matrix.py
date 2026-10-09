# -*- coding: utf-8 -*-
"""把京东两表 dump 成 JSON（供 node 侧生产引擎判定）"""
import io, json
from openpyxl import load_workbook

BASE = r'C:/Users/lzj/xwechat_files/wxid_auezz32wau7v22_5514/msg/file/2026-10/'
FILES = {
    'jd_order': '208386489_对账单下载_20260930_耙三样耙牛肉_7647722_7647723_0.xlsx',
    'jd_sku': '208386489_sku对账单下载_20260930_耙三样耙牛肉_7647724_7647725_0.xlsx',
}
OUT = r'C:/Users/lzj/WorkBuddy/Claw/catering-profit/_probe_tmp/_jd_matrix.json'

def norm(c):
    if c is None: return ''
    if hasattr(c, 'strftime'): return c.strftime('%Y-%m-%d %H:%M:%S')
    return c  # 保持原类型（数字/字符串），与 SheetJS 一致

res = {}
for key, fn in FILES.items():
    wb = load_workbook(BASE + fn, data_only=True, read_only=True)
    sheets = {}
    for sn in wb.sheetnames:
        rows = [[norm(c) for c in r] for r in wb[sn].iter_rows(values_only=True)]
        sheets[sn] = {'rows': rows}
    res[key] = {'file': fn, 'sheets': sheets}
    wb.close()

with io.open(OUT, 'w', encoding='utf-8') as f:
    json.dump(res, f, ensure_ascii=False)
print('写出', OUT, len(json.dumps(res)))
for k, v in res.items():
    for sn, s in v['sheets'].items():
        print(k, '|', sn, '|', len(s['rows']), '行 ×', max(len(r) for r in s['rows']) if s['rows'] else 0)
