# -*- coding: utf-8 -*-
"""B3 交叉验证：用**独立第三方实现**（Python openpyxl）读回 Node/SheetJS 生成的 xlsx。

为什么必须做这一步：
  tools/check_export_roundtrip.js 是「SheetJS 写 → SheetJS 读」= **同源往返**，
  只能证明「自己和自己一致」，证明不了「Excel / WPS / 别的库也认这份文件」。
  本脚本换一个**完全独立的实现**读同一份产物 ⇒ 才配得上「可回读性」的结论。

判据（与 Node 侧口径一致，但用 Decimal 精确回算，不经二进制浮点）：
  X1 工作表名与顺序        X2 利润表行数（含空行，空行也必须保留）
  X3 表头逐字              X4 金额列全为数字型（openpyxl data_type == 'n'）
  X5 利润表金额逐点回算分   X6 明细表行数        X7 明细表金额逐点回算
  X8 摊销来源行数          X9 摊销来源月均参考列逐点回算
  X10 口径说明零数字格      X11 金额均为「分」整数倍（无无限小数）
"""
import sys, os
from decimal import Decimal

sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook

XLSX_PATH = r'C:\Users\lzj\WorkBuddy\Claw\catering-profit\_probe_tmp\b3_export\月度报表.xlsx'

pass_n = 0
fail_n = 0
bad = []


def check(name, cond, detail=''):
    global pass_n, fail_n
    if cond:
        pass_n += 1
        print('✅ ' + name + (' · ' + detail if detail else ''))
    else:
        fail_n += 1
        bad.append(name)
        print('❌ ' + name + (' · ' + detail if detail else ''))


def fen_of(v):
    """元 → 分（Decimal 精确，不经二进制浮点）"""
    return int((Decimal(str(v)) * 100).to_integral_value())


print('文件 = ' + XLSX_PATH)
check('X0 产物存在且非空', os.path.exists(XLSX_PATH) and os.path.getsize(XLSX_PATH) > 1000,
      '%d 字节' % os.path.getsize(XLSX_PATH) if os.path.exists(XLSX_PATH) else '不存在')

wb = load_workbook(XLSX_PATH, data_only=True)
print('openpyxl 读到工作表 = ' + str(wb.sheetnames))

# ---------- 结构 ----------
check('X1 工作表名与顺序（利润表 / 收入费用明细 / 摊销来源 / 口径说明）',
      wb.sheetnames == ['利润表', '收入费用明细', '摊销来源', '口径说明'], str(wb.sheetnames))

ws = wb['利润表']
rows = [[c.value for c in r] for r in ws.iter_rows()]
check('X2 利润表行数 = 16（15 行内容 + 表头；含 2 个空行）', len(rows) == 16, '实得 %d' % len(rows))
check('X3 利润表表头逐字', rows[0] == ['项目', '金额(元)', '说明（这个数怎么来的）'], str(rows[0]))

# ---------- 类型（openpyxl: data_type 'n' = 数字） ----------
nonnum = []
for r in ws.iter_rows(min_row=2):
    c = r[1]
    if c.value is None or c.value == '':
        continue
    if c.data_type != 'n':
        nonnum.append('%s=%r(type=%s)' % (c.coordinate, c.value, c.data_type))
check('X4 利润表金额列全为数字型（会计 SUM 才不为 0）', not nonnum, '; '.join(nonnum[:5]) or '全数字')

# ---------- 值保真 ----------
EXP = [4000000, 1200000, 2200000, 0, 600000, 300000, 333333, -33333, 1500000, 633333]
got = []
for r in ws.iter_rows(min_row=2):
    if '毛利率' in str(r[0].value or ''):
        continue
    v = r[1].value
    if isinstance(v, (int, float)):
        got.append(fen_of(v))
check('X5 利润表金额逐点回算分 ≡ 源分（10 个点）', got == EXP, '实得 ' + str(got))

r2 = [[c.value for c in r] for r in wb['收入费用明细'].iter_rows()]
check('X6 明细表行数 = 7（表头 + 4 收入行 + 2 费用行）', len(r2) == 7, '实得 %d' % len(r2))
det = [fen_of(x[3]) for x in r2[1:] if isinstance(x[3], (int, float))]
check('X7 明细表金额逐点回算（大类 + 细项 + 合计）',
      det == [4000000, 2500000, 1500000, 4000000, 1200000, 1200000], str(det))

r3 = [[c.value for c in r] for r in wb['摊销来源'].iter_rows()]
check('X8 摊销来源行数 = 5（表头 + 2 笔资产 + 2 条合计）', len(r3) == 5, '实得 %d' % len(r3))
ast = [fen_of(x[6]) for x in r3[1:] if isinstance(x[6], (int, float))]
check('X9 摊销来源「月均参考」列逐点回算（装修 333333 / 小工具 80000 / 两条合计）',
      ast == [333333, 80000, 333333, 0], str(ast))

# ---------- 口径说明：零数字格 ----------
nums = []
for r in wb['口径说明'].iter_rows():
    for c in r:
        if c.data_type == 'n':
            nums.append(c.coordinate)
check('X10 口径说明零数字格（纯文本表）', not nums, ','.join(nums[:5]) or '零数字格')

# ---------- 无无限小数 ----------
allv = []
for sh, col in [('利润表', 1), ('收入费用明细', 3), ('摊销来源', 2), ('摊销来源', 6)]:
    for r in wb[sh].iter_rows(min_row=2):
        v = r[col].value
        if isinstance(v, (int, float)):
            allv.append((sh, v))
bad2 = ['%s=%r' % (s, v) for s, v in allv
        if (Decimal(str(v)) * 100) != (Decimal(str(v)) * 100).to_integral_value()]
check('X11 金额均为「分」整数倍（不得出现无限小数）', not bad2,
      '; '.join(bad2[:4]) or '共 %d 格全部 ≤2 位小数' % len(allv))

# ---------- 空行保留（读回时真存在，不是被裁掉） ----------
blank_ok = sum(1 for r in rows if all((c is None or c == '') for c in r))
check('X12 利润表 2 个空行在 openpyxl 读回时依然保留（少行不报错正是给会计的错账）',
      blank_ok == 2, '实得 %d 个空行' % blank_ok)

print('\n===== B3 交叉验证（openpyxl %s · 独立实现）：%d 通过 / %d 失败 ====='
      % (__import__('openpyxl').__version__, pass_n, fail_n))
for b in bad:
    print('   ❌ ' + b)
sys.exit(0 if fail_n == 0 else 1)
