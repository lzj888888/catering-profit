# -*- coding: utf-8 -*-
"""把两平台 2026-08 源账单 xlsx 转成 JSON 矩阵 fixture（离线自测用，零依赖）。

⚠️ read_only=False —— 项目红线：`openpyxl read_only=True 会少算 1 行`，必须用 ws.max_row 口径。
"""
import openpyxl, json, os, datetime

ROOT = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(ROOT, "fixtures")

SRC = [
    ("taobao_2026-08", "淘宝闪购.xlsx"),
    ("meituan_2026-08", "09bbd01042244a1fa98e0d0f1963211b.xlsx"),
]


def cell(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return v
    return str(v)


for tag, fn in SRC:
    p = os.path.join(FIX, fn)
    wb = openpyxl.load_workbook(p, data_only=True, read_only=False)
    out = {"_file": fn, "_tag": tag, "sheets": {}}
    for ws in wb.worksheets:
        rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        out["sheets"][ws.title] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "rows": rows,
        }
    dst = os.path.join(FIX, tag + ".matrix.json")
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    shape = [(s, v["max_row"], v["max_col"]) for s, v in out["sheets"].items()]
    print(tag, "->", os.path.basename(dst), os.path.getsize(dst), "bytes")
    print("   sheets(max_row,max_col):", shape)
