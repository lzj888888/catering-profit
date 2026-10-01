# -*- coding: utf-8 -*-
"""R183 · 账单列结构探针（纯读）

用途：判定已归档的两份外卖账单**是否含菜品名**，即是否满足 PLAN_2026-10-01
「阶段② 硬前置①：1~2 份真实收银/外卖『商品销售明细』导出表」。

用法：
    python review/evidence/r183_m331_gap/probe_r183_billcols.py [仓库根]
    （不带参数时按本文件位置自动推算仓库根）

🔴 铁律：openpyxl 必须 read_only=False —— read_only=True 会**少算 1 行**（旧坑）。
"""
import io
import os
import sys
import warnings

warnings.filterwarnings("ignore")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.abspath(os.path.join(HERE, "..", "..", ".."))
FIXDIR = os.path.join(ROOT, "review", "evidence", "r181l_stage1_import_feed", "fixtures")

FILES = [
    ("美团", "09bbd01042244a1fa98e0d0f1963211b.xlsx"),
    ("淘宝闪购", "淘宝闪购.xlsx"),
]

# 菜品级线索词：出现任一即说明可能含菜品维度
DISH_HINTS = ["菜品", "商品名称", "商品名", "菜名", "品名", "规格名称", "套餐名"]
QTY_HINTS = ["数量", "份数", "销售数", "销量"]
AMT_HINTS = ["金额", "总价", "单价"]


def scan(path):
    print("=" * 72)
    print("FILE:", path)
    if not os.path.exists(path):
        print("  !! 文件不存在 —— 无法判定")
        return None
    wb = load_workbook(path, read_only=False, data_only=True)
    print("SHEETS:", wb.sheetnames)
    verdict = {}
    for ws in wb.worksheets:
        rows, cols = ws.max_row, ws.max_column
        best_i, best_n, best_vals = 0, -1, []
        for i in range(1, min(rows, 5) + 1):
            vals = [ws.cell(row=i, column=c).value for c in range(1, cols + 1)]
            n = sum(1 for v in vals if v not in (None, ""))
            if n > best_n:
                best_i, best_n, best_vals = i, n, vals
        headers = [("" if v is None else str(v).strip()) for v in best_vals]
        headers = [h for h in headers if h != ""]
        dish = [h for h in headers if any(k in h for k in DISH_HINTS)]
        qty = [h for h in headers if any(k in h for k in QTY_HINTS)]
        amt = [h for h in headers if any(k in h for k in AMT_HINTS)]
        print("-" * 72)
        print("SHEET: %-16s max_row=%-6d max_col=%-4d 表头行=%d 表头列数=%d"
              % (ws.title, rows, cols, best_i, len(headers)))
        print("  菜品/名称类列 :", dish if dish else "（无）")
        print("  数量类列     :", qty if qty else "（无）")
        print("  金额/价格类列 :", amt[:12], ("..." if len(amt) > 12 else ""))
        print("  全部列名:", " | ".join(headers))
        verdict[ws.title] = {"rows": rows, "cols": cols, "header_row": best_i,
                             "n_header": len(headers), "dish_cols": dish}
    wb.close()
    return verdict


allv = {}
for tag, fn in FILES:
    print("\n\n########## %s ##########" % tag)
    allv[tag] = scan(os.path.join(FIXDIR, fn))

print("\n\n===== 结论 =====")
print("判据：商品侧若只有「金额 / 数量」而**无名称类列** ⇒ 属订单级账单，不含菜品级销量。")
for tag, v in allv.items():
    if not v:
        print("  %-6s : 无法判定" % tag)
        continue
    hits = [s for s, d in v.items() if d["dish_cols"]]
    print("  %-6s : 含名称类列的 sheet = %s ⇒ %s"
          % (tag, hits if hits else "无", "疑似含菜品维度" if hits else "❌ 不含菜品名"))
