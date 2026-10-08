# review/evidence/r241_next/probe_incoming.py
# R241 —— 「外部数据到达」预置复测脚本（第 1 步：把 xlsx 结构化 dump 成 JSON）
#
# 用途：李老师发来京东 / 美团账单或销量表后，一条命令把表的真实结构取出来，
#       供第 2 步 probe_incoming.js 喂给生产解析器做识别面复测。
#
# 纪律：本脚本**只读文件、只写本目录**，不改任何业务代码、不连数据库。
#
# 运行：
#   python probe_incoming.py <xlsx1> [xlsx2] ...
# 输出：
#   review/evidence/r241_next/incoming_dump.json
#
# ⚠️ 用 openpyxl 读，与云端 xlsx 库是两条不同的路（R181m 的教训：测试路径 ≠ 生产路径）。
#    本脚本只用来**看清真实列名**，识别面结论一律以第 2 步（生产解析器）为准。

import json
import os
import sys

try:
    import openpyxl
except ImportError:
    print('需要 openpyxl：pip install openpyxl', file=sys.stderr)
    sys.exit(2)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'incoming_dump.json')
PREVIEW_ROWS = 8   # 前 8 行足够覆盖「表头行位置不定」的情况


def cell(v):
    """归一成 JSON 可序列化的值；日期/时间戳保持原文以便人工核对。"""
    if v is None:
        return ''
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def dump_one(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= PREVIEW_ROWS:
                break
            rows.append([cell(x) for x in row])
        # 行数：read_only 下 max_row 可能不准，退化为扫描计数（表不大，可接受）
        n = 0
        for _ in ws.iter_rows(values_only=True):
            n += 1
        sheets.append({'name': name, 'rowCount': n, 'head': rows})
    wb.close()
    return {'file': os.path.basename(path), 'path': path, 'sheets': sheets}


def main():
    paths = sys.argv[1:]
    if not paths:
        print('用法：python probe_incoming.py <xlsx1> [xlsx2] ...', file=sys.stderr)
        sys.exit(1)
    out = []
    for p in paths:
        if not os.path.isfile(p):
            print('  [skip] 文件不存在：' + p, file=sys.stderr)
            continue
        try:
            d = dump_one(p)
            out.append(d)
            print('  [ok] %s：%d 个 sheet' % (d['file'], len(d['sheets'])))
            for s in d['sheets']:
                print('        · %-20s 行数=%-6d 首个非空表头=%r'
                      % (s['name'], s['rowCount'],
                         next((r for r in s['head'] if any(str(x).strip() for x in r)), [])))
        except Exception as e:
            print('  [err] %s：%s' % (p, e), file=sys.stderr)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print('\n已写出：' + OUT)
    print('下一步：node probe_incoming.js')


if __name__ == '__main__':
    main()
