# -*- coding: utf-8 -*-
"""生成 DSI-1 的两个样例 xlsx（给 AI 当 few-shot、给老板看形状）。

设计约束（必须遵守契约 specs/dev-specs/core/店算标准导入格式_v1.md）：
  · 第 1 行就是列名，不要标题行 / 不要两级表头 / 不要合并单元格 / 不要合计行
  · S1 只含菜品行（不含打包费）；销量允许小数（称重），退款填负数
"""
import warnings
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import os

OUT = os.path.join("specs", "dev-specs", "core", "examples")
os.makedirs(OUT, exist_ok=True)

def write(path, headers, rows, widths):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in rows:
        ws.append(r)
    wb.save(path)
    print("wrote", path, len(rows), "行")

# ---- S1 商品明细 ----
S1_H = ["日期", "渠道", "平台", "商品名称", "销量", "销售额", "备注"]
S1_R = [
    ["2026-09-07", "堂食", "", "重庆小面", 12, 216.00, ""],
    ["2026-09-07", "堂食", "", "卤牛肉", 1.5, 135.00, "称重"],
    ["2026-09-07", "外卖", "美团", "重庆小面", 5, 95.00, ""],
    ["2026-09-07", "外卖", "淘宝闪购", "重庆小面", 3, 57.00, ""],
    ["2026-09-08", "堂食", "", "重庆小面", 10, 180.00, ""],
    ["2026-09-08", "堂食", "", "卤牛肉", 2, 180.00, ""],
    ["2026-09-08", "外卖", "美团", "酸辣粉", 6, 108.00, ""],
    # 退款：填负数（不许删行）
    ["2026-09-08", "外卖", "美团", "重庆小面", -1, -19.00, "顾客退单"],
]
write(os.path.join(OUT, "dsi1_s1_example.xlsx"), S1_H, S1_R, [12, 8, 12, 16, 8, 10, 14])

# ---- S2 月渠道汇总 ----
S2_H = ["月份", "渠道", "平台", "营业额", "到手额", "订单数"]
S2_R = [
    ["2026-09", "堂食", "", 48600.00, 48600.00, 1620],
    ["2026-09", "外卖", "美团", 21300.00, 17700.00, 710],
    ["2026-09", "外卖", "淘宝闪购", 15800.00, 13900.00, 520],
]
write(os.path.join(OUT, "dsi1_s2_example.xlsx"), S2_H, S2_R, [10, 8, 12, 12, 12, 8])
