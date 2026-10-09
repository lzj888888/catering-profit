# -*- coding: utf-8 -*-
# review/evidence/r249_import/r249_recalc.py —— 独立复算锚点（不依赖仓内引擎，用 openpyxl 直读真文件）
# 用途：给 R249 真云 A/B 复验提供**独立**判据数字（口径：销量=销量列求和；金额=销售额列求和，单位元）
import io, sys
from openpyxl import load_workbook

P = r'C:/Users/lzj/WorkBuddy/Claw/catering-profit/_probe_tmp/taobao_goods.xlsx'
wb = load_workbook(P, data_only=True, read_only=True)
print('sheets =', wb.sheetnames)
ws = wb[wb.sheetnames[0]]

rows = list(ws.iter_rows(values_only=True))
print('总行数 =', len(rows))
hdr = [str(x).strip() if x is not None else '' for x in rows[0]]
print('表头 =', hdr)

def idx(name):
    return hdr.index(name)

i_date, i_name = idx('日期'), idx('商品名称')
i_qty, i_amt = idx('销量'), idx('销售额')

data = rows[1:]
data = [r for r in data if r and r[0] is not None]
print('数据行 =', len(data))

names = set()
dates = set()
sq = 0.0
sa = 0.0
q_pos = a_pos = zero_amt_q_pos = 0
for r in data:
    d = r[i_date]
    ds = d.strftime('%Y-%m-%d') if hasattr(d, 'strftime') else str(d)
    dates.add(ds)
    if r[i_name] is not None:
        names.add(str(r[i_name]).strip())
    q = float(r[i_qty] or 0)
    a = float(r[i_amt] or 0)
    sq += q
    sa += a
    if q > 0: q_pos += 1
    if a > 0: a_pos += 1
    if a == 0 and q > 0: zero_amt_q_pos += 1

print('唯一商品 =', len(names))
print('日期数 =', len(dates), sorted(dates))
print('Σqty =', sq)
print('Σamt(元) = %.2f' % sa)
print('Σamt(分) =', int(round(sa * 100)))
print('qty>0 行 =', q_pos, '| amt>0 行 =', a_pos, '| amt=0 且 qty>0 行 =', zero_amt_q_pos)

# 🔴 口径修正（R249 实测踩过）：真云 `zeroAmountQty` 是**按商品名聚合**后的条数
#   （见 service.js::parseDishSalesC 注释 v1.7 C-5：同名 SKU 多天各一行，按名合并为 1 条），
#   **不是原始行数**。第一次我拿"原始行数 7"去比"聚合条数 1"，误报 MISMATCH —— 是我的口径错，不是代码错。
zero_names = set()
for r in data:
    q = float(r[i_qty] or 0)
    a = float(r[i_amt] or 0)
    if a == 0 and q > 0 and r[i_name] is not None:
        zero_names.add(str(r[i_name]).strip())
print('amt=0 且 qty>0 的**按名聚合后**条数 =', len(zero_names), sorted(zero_names))

print()
print('--- 真云 importSalesBill 预览回执锚点（R249 实测）---')
print('groups=7 / totals.qty=118 / totals.amountFen=182308 / zeroAmountQtyLen=1')
print('本脚本 Σqty=%d  ΣamtFen=%d  zeroAmtQty(聚合)=%d'
      % (sq, int(round(sa * 100)), len(zero_names)))
ok = (int(sq) == 118 and int(round(sa * 100)) == 182308 and len(zero_names) == 1)
print('对齐结论 =', 'ALIGNED' if ok else 'MISMATCH')
sys.exit(0 if ok else 1)
