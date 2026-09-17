# -*- coding: utf-8 -*-
"""Markdown -> 纵向 A4 可打印 docx（表头每页重复 + 页码 + 引用框 + 代码框）。
支持： #/##/### 标题、- 列表、1. 编号行、| 表格、> 引用、```代码```、--- 分隔、
      行内 `code`、**粗体**、【粗体】。
用法： python tools/md2docx_portrait.py <in.md> <out.docx> <标题>
       ⚠️ 传参给 Windows python 一律用**显式 `C:/...` 绝对路径**，别用 shell 变量拼
          （`$R/x.py` 展开成 `/c/...` 会被当成 `c:\c\...` ⇒ `Errno 2`）。
"""
import sys, re
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Microsoft YaHei'
MONO = 'Consolas'
CODE_COLOR = RGBColor(0xB0, 0x30, 0x20)
HDR_FILL = '4472C4'
QUOTE_FILL = 'F2F2F2'
CODE_FILL = 'F7F7F7'


def set_font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), name)


def add_field(p, field):
    r = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = field
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1); r._r.append(it); r._r.append(f2)


def add_inline(p, text, size=10.5):
    for part in re.split(r'(`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('`') and part.endswith('`') and len(part) >= 2:
            r = p.add_run(part[1:-1]); r.font.size = Pt(size - 0.5)
            r.font.color.rgb = CODE_COLOR; set_font(r, MONO)
        else:
            for s in re.split(r'(\*\*[^*]+\*\*|【[^】]*】)', part):
                if not s:
                    continue
                if s.startswith('**') and s.endswith('**') and len(s) > 4:
                    r = p.add_run(s[2:-2]); r.bold = True; r.font.size = Pt(size); set_font(r)
                elif s.startswith('【') and s.endswith('】'):
                    r = p.add_run(s); r.bold = True; r.font.size = Pt(size); set_font(r)
                else:
                    r = p.add_run(s); r.font.size = Pt(size); set_font(r)


def shade_par(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true')
    trPr.append(th)


def add_code_block(doc, lines):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
    cell = t.cell(0, 0); cell.text = ''
    shade_cell(cell, CODE_FILL)
    for idx, ln in enumerate(lines):
        para = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        r = para.add_run(ln); r.font.size = Pt(9); set_font(r, MONO)
    doc.add_paragraph('')


def main():
    inp, outp, title = sys.argv[1], sys.argv[2], sys.argv[3]
    doc = Document()
    st = doc.styles['Normal']; st.font.size = Pt(10.5); st.font.name = FONT
    _rp = st.element.get_or_add_rPr(); _rf = _rp.get_or_add_rFonts()
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        _rf.set(qn(a), FONT)

    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)   # 纵向 A4
    sec.left_margin = Mm(18); sec.right_margin = Mm(18)
    sec.top_margin = Mm(16); sec.bottom_margin = Mm(16)

    hdr = sec.header.paragraphs[0]; hdr.text = ''
    hr = hdr.add_run(title); hr.bold = True; hr.font.size = Pt(11); set_font(hr)
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = hdr._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6'); bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), HDR_FILL)
    pbdr.append(bottom); pPr.append(pbdr)

    fp = sec.footer.paragraphs[0]; fp.text = ''; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run('第 '); r1.font.size = Pt(9); set_font(r1)
    add_field(fp, 'PAGE')
    r2 = fp.add_run(' 页 / 共 '); r2.font.size = Pt(9); set_font(r2)
    add_field(fp, 'NUMPAGES')
    r3 = fp.add_run(' 页'); r3.font.size = Pt(9); set_font(r3)

    lines = open(inp, encoding='utf-8').read().split('\n')
    i, n = 0, len(lines)
    while i < n:
        raw = lines[i].rstrip()
        if raw.strip() == '':
            i += 1; continue
        if raw.strip() == '---':
            doc.add_paragraph(''); i += 1; continue
        if raw.startswith('# '):
            p = doc.add_heading(level=0); p.text = ''; add_inline(p, raw[2:].strip(), 17)
            for r in p.runs: r.bold = True
            i += 1; continue
        if raw.startswith('## '):
            p = doc.add_heading(level=1); p.text = ''; add_inline(p, raw[3:].strip(), 13.5)
            for r in p.runs: r.bold = True
            i += 1; continue
        if raw.startswith('### '):
            p = doc.add_heading(level=2); p.text = ''; add_inline(p, raw[4:].strip(), 11.5)
            for r in p.runs: r.bold = True
            i += 1; continue
        if raw.startswith('```'):
            i += 1; buf = []
            while i < n and not lines[i].startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            add_code_block(doc, buf); continue
        if raw.startswith('>'):
            buf = []
            while i < n and lines[i].lstrip().startswith('>'):
                buf.append(lines[i].lstrip()[1:].strip()); i += 1
            p = doc.add_paragraph(); p.text = ''
            add_inline(p, ' '.join(buf), 10)
            pPr = p._p.get_or_add_pPr()
            ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '270'); pPr.append(ind)
            shade_par(p, QUOTE_FILL)
            continue
        if raw.startswith('|'):
            tl = []
            while i < n and lines[i].strip().startswith('|'):
                tl.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in r.strip('|').split('|')] for r in tl]
            data = []
            for idx, r in enumerate(rows):
                if idx == 1 and all(set(c) <= set('-: ') and c for c in r):
                    continue
                data.append(r)
            if data:
                ncol = max(len(r) for r in data)
                t = doc.add_table(rows=len(data), cols=ncol); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(data):
                    for ci in range(ncol):
                        c = t.cell(ri, ci); c.text = ''
                        para = c.paragraphs[0]
                        add_inline(para, row[ci] if ci < len(row) else '', 9.5)
                        if ri == 0:
                            shade_cell(c, HDR_FILL)
                            for r in para.runs:
                                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_repeat_header(t.rows[0])
                doc.add_paragraph('')
            continue
        if raw.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); p.text = ''
            add_inline(p, raw[2:].strip(), 10.5)
            i += 1; continue
        if raw[:2] in ('  ', '\t') and raw.strip():
            p = doc.add_paragraph(); p.text = ''
            add_inline(p, raw.strip(), 10)
            i += 1; continue
        p = doc.add_paragraph(); p.text = ''
        add_inline(p, raw.strip(), 10.5)
        i += 1

    doc.save(outp)
    print('SAVED', outp)


if __name__ == '__main__':
    main()
